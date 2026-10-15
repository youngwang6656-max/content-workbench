"""Bounded local extraction. Original files are never modified or executed."""
import csv,io,zipfile
from pathlib import Path

SUPPORTED={'.txt','.md','.pdf','.docx','.pptx','.xlsx','.csv','.tsv','.png','.jpg','.jpeg','.webp','.ppt','.doc','.xls'}


def extract(data,suffix):
    if suffix in ('.doc','.ppt','.xls'):
        return [{'label':'原文件','text':'','visual':False}],['旧版 Office 文件请另存为 DOCX、PPTX、XLSX 或 PDF，再导入读取。'],'needs_text'
    units=[];warnings=[]
    def add(label,text,visual=False,**extra):
        units.append(dict(label=label,text=text,visual=visual,**extra))
        if len(units)>200 or sum(len(u['text']) for u in units)>200000:raise ValueError('资料过大，请拆分为二百页/段、二十万字以内的文件')
    if suffix in ('.docx','.pptx','.xlsx'):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(x.file_size for x in z.infolist())>50_000_000:raise ValueError('文件展开后超过五十 MB，请拆分资料')
    if suffix in ('.txt','.md','.csv','.tsv'):
        try:text=data.decode('utf-8-sig')
        except UnicodeError:
            try:text=data.decode('gb18030')
            except UnicodeError:raise ValueError('文字编码无法识别，请另存为 UTF-8') from None
        if suffix in ('.csv','.tsv'):
            rows=list(csv.reader(io.StringIO(text),delimiter='\t' if suffix=='.tsv' else ','))
            text='\n'.join(f'第 {i+1} 行: '+' | '.join(row) for i,row in enumerate(rows))
        for i in range(0,max(1,len(text)),12000):add(f'文字段 {i//12000+1}',text[i:i+12000])
    elif suffix=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(data))
        if len(reader.pages)>200:raise ValueError('PDF 超过二百页，请拆分后导入')
        for i,page in enumerate(reader.pages):add(f'第 {i+1} 页',page.extract_text() or '',True,page=i)
        warnings.append('PDF 文字提取不能完整表达版式、图表和扫描内容；启用页面图片读取可辅助核对。')
    elif suffix=='.docx':
        from docx import Document
        doc=Document(io.BytesIO(data));parts=[]
        for block in doc.iter_inner_content():
            parts.append(block.text if hasattr(block,'text') else '\n'.join(' | '.join(c.text for c in row.cells) for row in block.rows))
        text='\n'.join(parts)
        # OOXML media names are references only, never filesystem destinations.
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            media=[x.filename for x in z.infolist() if x.filename.startswith('word/media/')]
            from lxml.etree import fromstring,XMLParser
            for name in z.namelist():
                if name.startswith(('word/header','word/footer','word/comments')) and name.endswith('.xml'):
                    root=fromstring(z.read(name),XMLParser(resolve_entities=False,no_network=True));extra=' '.join(e.text or '' for e in root.iter() if e.tag.endswith('}t'))
                    if extra.strip():text+='\n['+name+'] '+extra
            if b'txbxContent' in z.read('word/document.xml'):warnings.append('未读取：Word 文本框中的内容与布局，请另存 PDF 或补充文字。')
        for i in range(0,max(1,len(text)),12000):add(f'文档段 {i//12000+1}',text[i:i+12000])
        for i,name in enumerate(media):add(f'文档内图片 {i+1}','',True,media=name)
        warnings.append('DOCX 按正文/表格顺序提取，并附页眉、页脚和批注文字；文本框及分页版式请另存 PDF 核对。')
    elif suffix=='.pptx':
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        prs=Presentation(io.BytesIO(data))
        for i,slide in enumerate(prs.slides):
            texts=[];media=[]
            def visit(shapes):
                for shape in shapes:
                    if shape.shape_type==MSO_SHAPE_TYPE.GROUP:visit(shape.shapes)
                    if shape.has_text_frame:texts.append(shape.text)
                    if shape.has_table:texts.extend(' | '.join(c.text for c in row.cells) for row in shape.table.rows)
                    if shape.shape_type==MSO_SHAPE_TYPE.PICTURE:media.append(shape._element.blipFill.blip.rEmbed)
                    if shape.has_chart:
                        try:
                            if any(family in str(shape.chart.chart_type) for family in ('SCATTER','BUBBLE')):
                                warnings.append(f'未读取：第 {i+1} 张幻灯片含散点/气泡图，当前文字提取不保留完整坐标关系，请另存 PDF 或补充数据。')
                            texts.append('图表：'+' | '.join(str(c.label) for c in shape.chart.plots[0].categories))
                            texts.extend(str(s.name)+': '+str(list(s.values)) for s in shape.chart.series)
                        except (AttributeError,ValueError,TypeError):warnings.append(f'未读取：第 {i+1} 张幻灯片图表数据无法完整提取，请补充页面截图。')
            visit(slide.shapes)
            if 'drawingml/2006/diagram' in slide._element.xml: warnings.append(f'未读取：第 {i+1} 张幻灯片的 SmartArt，请另存 PDF 或补充截图。')
            if slide.has_notes_slide:texts.append('演讲者备注：'+slide.notes_slide.notes_text_frame.text)
            add(f'第 {i+1} 张幻灯片','\n'.join(texts),bool(media),slide=i,media=media)
        warnings.append('PPTX 提取文字、表格、图表数据与内嵌图片；不还原整页版式/SmartArt，布局关系请另存 PDF 或补充截图。')
    elif suffix=='.xlsx':
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if any(n.startswith(('xl/charts/','xl/media/','xl/drawings/')) for n in z.namelist()):warnings.append('未读取：Excel 内的图表、图片或绘图，请另存 PDF 或补充说明。')
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
        try:
            for sheet in book:
                if sheet.max_row>20000 or sheet.max_column>200:raise ValueError('表格超过两万行或二百列，请拆分后导入')
                lines=[];start=1
                for row in sheet.iter_rows():
                    text=' | '.join(f'{c.coordinate}={c.value}' for c in row if c.value is not None)
                    if text:lines.append(text)
                    if sum(map(len,lines))>=10000:
                        add(f'工作表 {sheet.title} · 行 {start}-{row[0].row}','\n'.join(lines));start=row[0].row+1;lines=[]
                if lines:add(f'工作表 {sheet.title} · 行 {start}-{sheet.max_row}','\n'.join(lines))
        finally:book.close()
        warnings.append('按单元格读取，公式保留为公式且不重新计算；图表、图片与样式未读取，请另存 PDF 或补充截图。')
    else:
        if not (data.startswith(b'\x89PNG\r\n\x1a\n') or data.startswith(b'\xff\xd8\xff') or data[:4]==b'RIFF' and data[8:12]==b'WEBP'):raise ValueError('图片格式不正确')
        add('图片 1','',True)
        return units,warnings,'needs_image'
    if not units:add('文档 1','')
    state='text' if all(u['text'].strip() for u in units) and not any(w.startswith('未读取：') for w in warnings) else 'needs_text'
    return units,warnings,state


def images_for(path,units,outdir):
    """Render selected PDF pages or embedded pictures only, with strict bounds."""
    from PIL import Image
    result=[]
    def save(image):
        if len(result)>=8:raise ValueError('本次包含超过八张图片，请缩小读取范围')
        image.thumbnail((1600,1600));target=Path(outdir)/f'{len(result)+1}.png'
        image.convert('RGB').save(target);result.append(target)
    suffix=path.suffix.lower()
    if suffix=='.pdf':
        import pypdfium2 as pdfium
        with pdfium.PdfDocument(str(path)) as doc:
            for unit in units:
                page=doc[unit['page']]
                try:
                    scale=min(2,1600/max(page.get_size()))
                    bitmap=page.render(scale=scale)
                    try:
                        img=bitmap.to_pil()
                        try:save(img)
                        finally:img.close()
                    finally:bitmap.close()
                finally:page.close()
    elif suffix=='.pptx':
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        prs=Presentation(str(path))
        for unit in units:
            slide=prs.slides[unit['slide']]
            width,height=prs.slide_width,prs.slide_height
            scale=min(1600/width,1600/height)
            canvas=Image.new('RGB',(max(1,int(width*scale)),max(1,int(height*scale))),'white')
            found=False
            def place(shapes):
                nonlocal found
                for shape in shapes:
                    if shape.shape_type==MSO_SHAPE_TYPE.GROUP:
                        place(shape.shapes)
                    elif shape.shape_type==MSO_SHAPE_TYPE.PICTURE:
                        found=True
                        with Image.open(io.BytesIO(shape.image.blob)) as picture:
                            picture=picture.convert('RGB')
                            picture.thumbnail((max(1,int(shape.width*scale)),max(1,int(shape.height*scale))))
                            canvas.paste(picture,(max(0,int(shape.left*scale)),max(0,int(shape.top*scale))))
            place(slide.shapes)
            if found:save(canvas)
    elif suffix=='.docx':
        with zipfile.ZipFile(path) as z:
            for unit in units:
                if unit.get('media'):
                    with Image.open(io.BytesIO(z.read(unit['media']))) as img:save(img)
    else:
        with Image.open(path) as img:save(img)
    return result
