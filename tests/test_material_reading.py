import asyncio,io,json
import pytest
from test_repository import make
from waypoint.materials import Materials


def test_office_extracts_positions_and_keeps_original(tmp_path):
    from pptx import Presentation
    from openpyxl import Workbook
    r,p,t=make(tmp_path); service=Materials(r)
    slides=Presentation();slide=slides.slides.add_slide(slides.slide_layouts[1]);slide.shapes.title.text='产品参数';slide.placeholders[1].text='续航 12 小时'
    stream=io.BytesIO();slides.save(stream)
    ppt=service.import_bytes(p,'产品.pptx',stream.getvalue())
    assert '第 1 张' in ppt['text'] and '12 小时' in ppt['text']
    assert (r.db.root/'attachments'/ppt['path']).read_bytes()==stream.getvalue()
    book=Workbook();book.active.title='参数';book.active.append(['重量',42]);stream=io.BytesIO();book.save(stream)
    sheet=service.import_bytes(p,'规格.xlsx',stream.getvalue())
    assert 'B1=42' in sheet['text'] and '参数' in sheet['text']


def test_analysis_adoption_and_supplement_preserve_source(tmp_path):
    from waypoint.material_reading import analyze,adopt,supplement,context_material
    from waypoint.workflow import Engine
    from test_workflow import FakeProvider
    r,p,t=make(tmp_path);r.save_settings({'model_roles':{'write':'test','read':'reader'}})
    m=Materials(r).add_text(p,'规格','Original specification 42.')
    class Reader(FakeProvider):
        async def capabilities(self):return {'auth_state':'chatgpt','models':[{'model':'test'},{'model':'reader'}]}
        async def generate(self,request):
            self.requests.append(request)
            yield {'kind':'completed','text':json.dumps({'summary_zh':'资料摘要','facts_zh':['[文字段 1] 参数为 42'],'missing_zh':[],
                'unit_notes_zh':[{'number':1,'note_zh':'规格注明参数为 42'}]})}
    reader=Reader();engine=Engine(r,reader)
    result=asyncio.run(analyze(engine,m['id'],{'start':1,'end':1,'include_images':False}))
    assert reader.requests[0]['model']=='reader'
    assert reader.requests[0]['effort']=='high'
    assert r.get(m['id'],'material')['text']==m['text']
    assert not context_material(r.get(m['id'],'material')).get('adopted_reading')
    adopt(r,m['id'],result['id'])
    assert '参数为 42' in context_material(r.get(m['id'],'material'))['text']
    assert '规格注明参数为 42' in context_material(r.get(m['id'],'material'))['text']
    supplement(r,m['id'],'用户补充：单位为克',True)
    assert r.get(m['id'],'material')['text']==m['text']
    assert len(r.get(m['id'],'material')['analysis_history'])==1
    with pytest.raises(ValueError,match='变化'):adopt(r,m['id'],result['id'])


def test_material_failure_has_exact_jump_target(tmp_path):
    from test_workflow import setup
    r,p,t,f,e=setup(tmp_path)
    m=Materials(r).add_text(None,'扫描件','',extraction='needs_text')
    from waypoint.contracts import TaskSpec
    t=r.create_task(TaskSpec(project_id=p,kind='cb',raw_input='依据资料写 CB'))
    job=asyncio.run(e.run_and_wait(t))
    assert job['status']=='paused'
    assert r.task(t)['attention']=={'page':'global-materials','material_id':m['id'],'project_id':None}
    assert f.requests==[]


def test_partial_document_cannot_claim_full_reading(tmp_path):
    from waypoint.material_reading import analyze,adopt,context_material
    from waypoint.workflow import Engine
    from test_workflow import FakeProvider
    r,p,t=make(tmp_path);r.save_settings({'model_roles':{'write':'test'}})
    m=Materials(r).add_text(p,'长资料','',extraction='needs_text',units=[{'label':f'第 {i} 页','text':'','visual':False} for i in range(1,10)])
    class Reader(FakeProvider):
        async def generate(self,request):
            yield {'kind':'completed','text':json.dumps({'summary_zh':'空白页面','facts_zh':[],'missing_zh':['没有可读取文字']})}
    with pytest.raises(ValueError):asyncio.run(analyze(Engine(r,Reader()),m['id'],{'start':1,'end':9}))
    assert not r.get(m['id'],'material').get('analysis_history')


def test_url_only_requires_real_web_evidence(tmp_path):
    from waypoint.material_reading import analyze
    from waypoint.workflow import Engine
    from test_workflow import FakeProvider
    r,p,t=make(tmp_path);r.save_settings({'model_roles':{'write':'test'}})
    m=Materials(r).add_text(p,'网址','',url='https://example.com/spec')
    class Reader(FakeProvider):
        supports_web_research=True
        async def generate(self,request):
            yield {'kind':'completed','text':json.dumps({'summary_zh':'没有取证','facts_zh':[],'missing_zh':[]})}
    with pytest.raises(ValueError,match='证据'):asyncio.run(analyze(Engine(r,Reader()),m['id'],{}))
    assert not r.get(m['id'],'material').get('analysis_history')


def test_pdf_and_embedded_images_render_without_losing_locations(tmp_path):
    from waypoint.document_text import extract,images_for
    from PIL import Image
    from pypdf import PdfWriter
    from pptx import Presentation
    from pptx.util import Inches
    from docx import Document
    image=io.BytesIO();Image.new('RGB',(30,20),'blue').save(image,format='PNG')
    writer=PdfWriter();writer.add_blank_page(200,100);buffer=io.BytesIO();writer.write(buffer)
    cases=[('.pdf',buffer.getvalue())]
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.shapes.add_picture(io.BytesIO(image.getvalue()),0,0,Inches(1));buffer=io.BytesIO();prs.save(buffer);cases.append(('.pptx',buffer.getvalue()))
    doc=Document();doc.add_paragraph('图片说明');doc.add_picture(io.BytesIO(image.getvalue()));buffer=io.BytesIO();doc.save(buffer);cases.append(('.docx',buffer.getvalue()))
    for ext,data in cases:
        path=tmp_path/('source'+ext);path.write_bytes(data);units,warnings,state=extract(data,ext)
        assert state=='needs_text'
        folder=tmp_path/ext[1:];folder.mkdir()
        paths=images_for(path,[u for u in units if u['visual']],folder)
        assert len(paths)==1 and Image.open(paths[0]).size[0]<=1600


def test_missing_original_never_becomes_complete_from_excerpt(tmp_path):
    from waypoint.material_reading import analyze,adopt,context_material
    from waypoint.workflow import Engine
    from test_workflow import FakeProvider
    r,p,t=make(tmp_path);r.save_settings({'model_roles':{'write':'test'}})
    m=Materials(r).add_text(p,'缺失原件','只有摘录',missing_attachment=True,extraction='needs_text')
    class Reader(FakeProvider):
        async def generate(self,request):yield {'kind':'completed','text':json.dumps({'summary_zh':'摘录摘要','facts_zh':[],'missing_zh':[]})}
    result=asyncio.run(analyze(Engine(r,Reader()),m['id'],{'include_images':False}))
    assert not result['complete'] and '原附件' in result['missing_zh'][0]
    adopt(r,m['id'],result['id'])
    assert context_material(r.get(m['id'],'material'))['extraction']=='needs_text'


def test_headers_and_known_gaps_not_silently_lost(tmp_path):
    from waypoint.document_text import extract
    from docx import Document
    from openpyxl import Workbook
    from openpyxl.chart import BarChart,Reference
    doc=Document();doc.add_paragraph('Description');doc.sections[0].header.paragraphs[0].text='Indoor use only'
    stream=io.BytesIO();doc.save(stream);units,_,_=extract(stream.getvalue(),'.docx')
    assert 'Indoor use only' in '\n'.join(u['text'] for u in units)
    book=Workbook();book.active.append([1,2]);chart=BarChart();chart.add_data(Reference(book.active,min_col=1,max_col=2,min_row=1));book.active.add_chart(chart,'D1')
    stream=io.BytesIO();book.save(stream);units,warnings,state=extract(stream.getvalue(),'.xlsx')
    assert state=='needs_text' and any(w.startswith('未读取：') for w in warnings)


def test_accepted_long_document_uses_bounded_analysis(tmp_path):
    from waypoint.material_reading import fingerprint,adopt
    from waypoint.context import build_context
    r,p,t=make(tmp_path)
    m=Materials(r).add_text(p,'长资料','原文'*40000)
    m['analysis_history']=[{'id':'analysis1','source_fingerprint':fingerprint(m),'start':1,'end':7,'complete':True,'summary_zh':'摘要','facts_zh':['[段落 1] 产品条件'],'missing_zh':[],
        'unit_notes_zh':[{'number':i,'note_zh':f'第 {i} 段条件'} for i in range(1,8)]}]
    m['original_text_history']=[{'text':'以前的原文'*30000}]
    r.put('material',m);adopt(r,m['id'],'analysis1')
    c=build_context(r,t,'english');assert len(c['materials'][0]['text'])<1000
    assert '第 7 段条件' in c['materials'][0]['text']
    assert len(json.dumps(c['materials']))<2000
    assert r.get(m['id'],'material')['text']=='原文'*40000


def test_old_analysis_without_page_evidence_remains_due_for_full_reading(tmp_path):
    from waypoint.material_reading import describe,fingerprint
    r,p,t=make(tmp_path)
    units=[{'label':f'第 {i} 页','text':f'第 {i} 页内容'} for i in range(1,11)]
    m=Materials(r).add_text(p,'旧分析资料','原文',units=units)
    m['analysis_history']=[{'id':'old','source_fingerprint':fingerprint(m),'start':1,'end':8,'complete':True,
        'summary_zh':'旧版摘要','facts_zh':[],'missing_zh':[]}]
    m['adopted_analysis_ids']=['old']
    assert describe(m)['remaining']==list(range(1,11))


def test_attention_survives_project_transfer(tmp_path):
    from waypoint.transfer import ProjectTransfer
    r,p,t=make(tmp_path);m=Materials(r).add_text(p,'资料','来源文字')
    r.set_task_state(t,status='paused',attention={'page':'global-materials','material_id':m['id'],'project_id':p})
    service=ProjectTransfer(r);data=service.export(p);preview=service.preview([('project.zip',data)])
    imported=service.commit(preview['token'])['project_id']
    task=r.list('task',project_id=imported)[0];material=r.list('material',project_id=imported)[0]
    assert task['attention']=={'page':'materials','material_id':material['id'],'project_id':imported}


def test_scatter_chart_coordinates_require_visual_verification():
    from pptx import Presentation
    from pptx.chart.data import XyChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from waypoint.document_text import extract
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6]);data=XyChartData();series=data.add_series('Measured');series.add_data_point(1,42);series.add_data_point(2,65)
    slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER,0,0,1000000,1000000,data)
    stream=io.BytesIO();prs.save(stream);units,warnings,state=extract(stream.getvalue(),'.pptx')
    assert state=='needs_text' and any(w.startswith('未读取：') for w in warnings)


def test_long_uploaded_material_is_usable_without_ai_adoption(tmp_path):
    from waypoint.context import build_context
    r,p,t=make(tmp_path)
    units=[{'label':f'第 {i} 页','text':('普通说明 '*600 if i<38 else '特殊兼容限制：仅适用于车型 Z9。')} for i in range(1,39)]
    m=Materials(r).add_text(p,'产品说明','\n'.join(u['text'] for u in units),kind='attachment',units=units,extraction='needs_text')
    task=r.task(t);task['source_ids']=[m['id']];task['raw_input']='请讨论车型 Z9 的兼容限制';r.put('task',task)
    context=build_context(r,t,'english')
    selected=context['materials'][0]
    assert '第 38 页' in selected['text'] and '车型 Z9' in selected['text']
    assert selected['extraction']=='text'
    assert selected['source_extraction']=='needs_text'
    assert selected['local_indexed_units']==38
    assert len(selected['text'])<=12000
    assert not r.get(m['id'],'material').get('adopted_analysis_ids')


def test_pptx_many_embedded_images_becomes_one_bounded_slide_image(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image
    from waypoint.document_text import extract,images_for
    image=io.BytesIO();Image.new('RGB',(30,20),'blue').save(image,format='PNG')
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    for i in range(10):slide.shapes.add_picture(io.BytesIO(image.getvalue()),Inches(i%5),Inches(i//5),Inches(.4))
    buffer=io.BytesIO();prs.save(buffer);path=tmp_path/'many.pptx';path.write_bytes(buffer.getvalue())
    units,_,_=extract(buffer.getvalue(),'.pptx')
    outdir=tmp_path/'rendered';outdir.mkdir();images=images_for(path,units,outdir)
    assert len(images)==1
    with Image.open(images[0]) as visual:assert max(visual.size)<=1600


def test_brand_attachments_are_searchable_without_manual_selection(tmp_path):
    from waypoint.context import build_context
    r,p,t=make(tmp_path)
    own=Materials(r).add_text(p,'本品牌附件','车型 Z9 仅支持蓝牙连接',kind='attachment')
    other_project=r.create_project('其他品牌')
    other=Materials(r).add_text(other_project,'其他品牌附件','车型 Z9 支持有线连接',kind='attachment')
    global_attachment=Materials(r).add_text(None,'全局原始附件','车型 Z9 支持蜂窝连接',kind='attachment')
    task=r.task(t);task['raw_input']='说明车型 Z9 连接方式';r.put('task',task)
    found={m['id'] for m in build_context(r,t,'english')['materials']}
    assert own['id'] in found
    assert other['id'] not in found and global_attachment['id'] not in found
